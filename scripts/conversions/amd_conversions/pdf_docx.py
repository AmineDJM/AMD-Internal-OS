"""
PDF → WORD HAUTE FIDÉLITÉ (décision 0019).

Moteur : pdf2docx (PyMuPDF) — il reconstruit la mise en page (paragraphes, tableaux par leurs traits, images à leur
place et à leur taille, colonnes, polices, tailles, couleurs, gras/italique). Mesuré sur le corpus contre LibreOffice
(`--infilter=writer_pdf_import` : des cadres de dessin, pas un texte éditable) et une composition maison (texte seul) :
seul pdf2docx garde images ET tableaux ET un texte qui se reflue (`docs/assistant/decisions/0019-…`).

Ce que nous ajoutons, là où pdf2docx est faible :
  1. AVANT la reconstruction : le texte de chaque page est remis en ORDRE LOGIQUE et normalisé (`texte_pdf.py`) ;
     les en-têtes et pieds répétés sont repérés et retirés du corps ;
  2. APRÈS : le Word est relu (`docx_post.py`) — lignes recousues, listes, w:bidi/w:rtl/w:cs, tableaux RTL ;
     en-têtes et pieds reposés dans les VRAIES zones d'en-tête et de pied (numéro de page en champ PAGE) ;
  3. les pages SCANNÉES (sans texte, ou dont le texte est illisible) passent par l'OCR (`ocr.py`) : tesseract
     ara+fra+eng recompose une page texte à la même place, les illustrations découpées gardées en images ; si la
     page ne se reconstruit pas (confiance trop basse), son IMAGE est gardée et le texte reconnu suit.
"""

from __future__ import annotations

import logging
import re
import tempfile
from collections import Counter
from pathlib import Path

import pymupdf

from . import docx_post
from .normaliser import est_arabe
from .texte_pdf import Statistiques, remettre_a_l_endroit, texte_des_lignes

FLAGS = pymupdf.TEXT_MEDIABOX_CLIP | pymupdf.TEXT_CID_FOR_UNKNOWN_UNICODE
#: Une page dont plus de 5 % des caractères sont inconnus (police sans ToUnicode) se lit mieux par l'OCR.
SEUIL_INCONNUS = 0.05


def _brut(page) -> list:
    flags = FLAGS | (pymupdf.TEXT_INHIBIT_SPACES if _espaces_reelles(page) else 0)
    return completer_par_les_glyphes(page, page.get_text("rawdict", flags=flags)["blocks"])


def completer_par_les_glyphes(page, blocs: list) -> list:
    """Une grappe arabe couverte par un `/ActualText` peut DISPARAÎTRE de la lecture de PyMuPDF (mesuré : le « ضِ » final
    d'une ligne en Amiri). Relue glyphe par glyphe (`TEXT_IGNORE_ACTUALTEXT`), elle est là. On réinjecte chaque lettre
    arabe que la première lecture ne couvre pas (et ses marques), dans la ligne de même hauteur ; rien d'autre ne change."""
    from .normaliser import classe

    principaux = [ch for b in blocs if b.get("type", 0) == 0 for l in b["lines"] for s in l["spans"] for ch in s["chars"]]
    if not any(classe(ch["c"]) == "R" for ch in principaux):
        return blocs
    glyphes = [ch for b in page.get_text("rawdict", flags=FLAGS | pymupdf.TEXT_IGNORE_ACTUALTEXT | pymupdf.TEXT_INHIBIT_SPACES)["blocks"]
               if b.get("type", 0) == 0 for l in b["lines"] for s in l["spans"] for ch in s["chars"]]

    def meme_rangee(a, b) -> bool:
        return min(a[3], b[3]) - max(a[1], b[1]) > 0.4 * min(a[3] - a[1], b[3] - b[1])

    def couvert(g) -> bool:
        gx0, _, gx1, _ = g["bbox"]
        for m in principaux:
            if classe(m["c"]) != "R" or not meme_rangee(m["bbox"], g["bbox"]):
                continue
            mx0, _, mx1, _ = m["bbox"]
            if mx1 - mx0 > 0.05 and min(gx1, mx1) - max(gx0, mx0) > 0.3 * (gx1 - gx0):
                return True
            if mx1 - mx0 <= 0.05 and m["c"] == g["c"] and gx0 - 0.5 <= mx0 <= gx1 + 0.5:
                return True
        return False

    def hors_etendue(g) -> bool:
        """Au-delà des lettres arabes de sa rangée (une grappe perdue est en FIN de ligne, mesuré) — à l'intérieur, les
        boîtes des deux lectures diffèrent trop (Amiri) pour décider sans risque."""
        voisines = [m for m in principaux if classe(m["c"]) == "R" and meme_rangee(m["bbox"], g["bbox"])]
        if not voisines:
            return False
        # Bout de ligne arabe = à GAUCHE (sa fin logique) ; à droite, l'écho d'une lettre déjà lue se confondrait.
        return g["bbox"][2] <= min(m["bbox"][0] for m in voisines) + 0.5

    manquantes = [g for g in glyphes if classe(g["c"]) == "R" and g["bbox"][2] - g["bbox"][0] > 0.05 and hors_etendue(g) and not couvert(g)]
    for g in manquantes:
        ligne = next((l for b in blocs if b.get("type", 0) == 0 for l in b["lines"] if meme_rangee(tuple(l["bbox"]), g["bbox"])
                      and l["bbox"][0] - 3 * (g["bbox"][2] - g["bbox"][0] + 1) <= g["bbox"][0] <= l["bbox"][2] + 30), None)
        if ligne is None or not ligne["spans"]:
            continue
        marques = [m for m in glyphes if classe(m["c"]) == "M" and g["bbox"][0] - 0.5 <= m["bbox"][0] <= g["bbox"][2] + 0.5
                   and meme_rangee(m["bbox"], g["bbox"])]
        span = min(ligne["spans"], key=lambda s: abs(s["bbox"][0] - g["bbox"][0]))
        span["chars"].extend([{**g}, *({**m} for m in marques)])
        ligne["bbox"] = (min(ligne["bbox"][0], g["bbox"][0]), ligne["bbox"][1], max(ligne["bbox"][2], g["bbox"][2]), ligne["bbox"][3])
    return blocs


# ─────────────────────────── En-têtes et pieds ───────────────────────────

def reperer_entetes_pieds(doc: pymupdf.Document) -> dict:
    """Une ligne posée dans la marge haute (ou basse) de la majorité des pages, au même texte (chiffres masqués : un
    numéro de page varie) = un en-tête (un pied). Rend, par page, les zones à retirer du corps, et le texte à reposer."""
    if doc.page_count < 2:
        return {"zones": {}, "entete": [], "pied": []}
    vues: Counter = Counter()
    par_page: dict[int, list] = {}
    for page in doc:
        h = page.rect.height
        lignes = texte_des_lignes(remettre_a_l_endroit(_brut(page)))
        candidates = []
        for bbox, texte, rtl in lignes:
            if not texte.strip():
                continue
            zone = "entete" if bbox[3] < h * 0.12 else "pied" if bbox[1] > h * 0.88 else None
            if zone:
                cle = (zone, re.sub(r"\d+", "#", texte.strip()))
                candidates.append((cle, bbox, texte.strip(), rtl, page.rect.width))
        par_page[page.number] = candidates
        vues.update({c[0] for c in candidates})
    seuil = max(2, int(0.6 * doc.page_count + 0.5))
    retenues = {k for k, n in vues.items() if n >= seuil}
    zones: dict[int, list] = {}
    entete, pied = [], []
    for num, cands in par_page.items():
        for cle, bbox, texte, rtl, largeur in cands:
            if cle in retenues:
                zones.setdefault(num, []).append((bbox[0] - 1, bbox[1] - 1, bbox[2] + 1, bbox[3] + 1))
                if num == 0 or not any(e["cle"] == cle for e in (entete + pied)):
                    centre = (bbox[0] + bbox[2]) / 2
                    align = "center" if abs(centre - largeur / 2) < largeur * 0.08 else "right" if centre > largeur / 2 else "left"
                    ligne = {"cle": cle, "texte": texte, "rtl": rtl, "align": align, "numero": "#" in cle[1] and texte != cle[1], "y": bbox[1]}
                    if not any(e["cle"] == cle for e in (entete + pied)):
                        (entete if cle[0] == "entete" else pied).append(ligne)
    return {"zones": zones, "entete": sorted(entete, key=lambda e: e["y"]), "pied": sorted(pied, key=lambda e: e["y"])}


def _poser_zone(zone, lignes: list[dict]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    zone.is_linked_to_previous = False
    premier = True
    for l in lignes:
        p = zone.paragraphs[0] if premier and zone.paragraphs else zone.add_paragraph()
        premier = False
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "left": WD_ALIGN_PARAGRAPH.LEFT}[l["align"]]
        if l["numero"]:
            # Le numéro de page devient un champ PAGE : il se recalcule dans Word.
            avant, _, apres = re.split(r"(\d+)", l["texte"], maxsplit=1) if re.search(r"\d+", l["texte"]) else (l["texte"], "", "")
            if avant:
                p.add_run(avant)
            r = p.add_run()
            for typ, txt in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
                if typ:
                    f = OxmlElement("w:fldChar")
                    f.set(qn("w:fldCharType"), typ)
                    r._r.append(f)
                elif txt.strip() == "PAGE":
                    it = OxmlElement("w:instrText")
                    it.set(qn("xml:space"), "preserve")
                    it.text = txt
                    r._r.append(it)
                else:
                    t = OxmlElement("w:t")
                    t.text = txt
                    r._r.append(t)
            if apres:
                p.add_run(apres)
        else:
            p.add_run(l["texte"])
        r0 = p.runs[0] if p.runs else None
        if r0 is not None:
            r0.font.size = None


# ─────────────────────────── La conversion ───────────────────────────

def _espaces_reelles(page) -> bool:
    """Le PDF porte-t-il ses espaces (glyphes) ? Si oui, on interdit à PyMuPDF d'en SYNTHÉTISER : il en invente au
    milieu des mots arabes vocalisés (boîtes nulles de LibreOffice). Sinon (pdfTeX…), on garde la synthèse."""
    sans = page.get_text("text", flags=FLAGS | pymupdf.TEXT_INHIBIT_SPACES)
    lettres = sum(1 for c in sans if c.isalpha())
    # Un texte courant porte une espace toutes les 5 à 8 lettres ; sous 1 pour 16, le PDF n'en pose pas lui-même.
    return lettres < 20 or sans.count(" ") / lettres > 0.06


def _premier_mot(lignes, rtl: bool) -> float:
    """Largeur du premier mot (logique) d'une rangée : s'il tenait dans la place libre au bout de la rangée d'avant,
    le retour était voulu (fin de paragraphe, adresse…) ; sinon c'est un retour à la ligne automatique."""
    chars = [c for l in lignes for s in l.spans for c in getattr(s, "chars", [])]
    if not chars:
        return 0.0
    chars.sort(key=lambda c: c.bbox[0], reverse=rtl)
    bord = chars[0].bbox[2] if rtl else chars[0].bbox[0]
    for c in chars:
        if not c.c.strip():
            break
        bord = min(bord, c.bbox[0]) if rtl else max(bord, c.bbox[2])
    return abs(bord - (chars[0].bbox[2] if rtl else chars[0].bbox[0]))


class _Patch:
    """Branche nos correctifs dans pdf2docx, le temps d'une conversion :
      • extraction du texte : ordre logique, normalisation, en-têtes retirés, espaces synthétiques évitées ;
      • rangement des lignes d'une rangée : de droite à gauche quand elle est arabe ;
      • retours à la ligne : un retour n'est gardé que s'il est VOULU (le mot suivant tenait dans la place libre) ;
        sinon les lignes se recousent par une espace et le texte se reflue dans Word."""

    def __init__(self, zones: dict, stats: Statistiques):
        self.zones = zones
        self.stats = stats
        self.originaux: list = []

    def _remplacer(self, cible, nom: str, fn) -> None:
        self.originaux.append((cible, nom, getattr(cible, nom)))
        setattr(cible, nom, fn)

    def __enter__(self):
        from pdf2docx.common.Collection import Collection
        from pdf2docx.common.share import TextAlignment
        from pdf2docx.page import RawPageFitz as module
        from pdf2docx.text.Char import Char
        from pdf2docx.text.Lines import Lines

        zones, stats = self.zones, self.stats
        origine_texte = module.RawPageFitz._preprocess_text
        origine_tri = Collection.sort_in_line_order
        origine_retours = Lines.parse_line_break

        def _preprocess_text(page_self, **settings):
            page = page_self.page_engine
            if _espaces_reelles(page):
                brut = page.get_text
                page.get_text = lambda *a, **k: brut(*a, **{**k, "flags": k.get("flags", 0) | pymupdf.TEXT_INHIBIT_SPACES}) if a and a[0] == "rawdict" else brut(*a, **k)
                try:
                    blocs = origine_texte(page_self, **settings)
                finally:
                    del page.get_text
            else:
                blocs = origine_texte(page_self, **settings)
            return remettre_a_l_endroit(completer_par_les_glyphes(page, blocs), stats, exclure=zones.get(page.number))

        def sort_in_line_order(coll):
            origine_tri(coll)
            if not coll.is_vertical_text and len(coll) > 1:
                texte = "".join(getattr(e, "text", "") or "" for e in coll)
                if est_arabe(texte):
                    coll._instances.sort(key=lambda e: (-e.bbox.x1, e.bbox.y0))
            return coll

        def parse_line_break(lines_self, bbox, line_break_width_ratio, line_break_free_space_ratio):
            block = lines_self.parent
            if not block.is_horizontal_text:
                return origine_retours(lines_self, bbox, line_break_width_ratio, line_break_free_space_ratio)
            rangees = lines_self.group_by_physical_rows()
            for i, lignes in enumerate(rangees):
                lignes.sort_in_line_order()
                for l in lignes:
                    l.line_break = 0
                if i == len(rangees) - 1:
                    continue
                texte = "".join(l.text for l in lignes)
                rtl = est_arabe(texte)
                fin = lignes[-1]  # rangée déjà triée dans l'ordre logique (arabe : de droite à gauche)
                # La place libre se mesure jusqu'au bord de la ZONE (page, colonne, cellule) — pas du bloc, que la ligne
                # la plus longue définit : un titre suivi d'un sous-titre plus court ne serait jamais « voulu ».
                # Un bloc de deux rangées a pour largeur sa ligne la plus longue : on mesure alors jusqu'au bord de la zone.
                if len(rangees) <= 2:
                    # La zone = l'étendue des blocs de texte voisins (la colonne, la page sans ses marges, la cellule).
                    parent = block.parent
                    voisins = [b for b in getattr(parent, "_instances", []) if getattr(b, "is_text_block", False)] or [block]
                    gauche = min(b.bbox[0] for b in voisins)
                    droite = max(b.bbox[2] for b in voisins)
                else:
                    gauche, droite = block.bbox[0], block.bbox[2]
                if block.alignment == TextAlignment.CENTER:
                    libre = (droite - gauche) - (max(l.bbox[2] for l in lignes) - min(l.bbox[0] for l in lignes))
                elif rtl:
                    libre = min(l.bbox[0] for l in lignes) - gauche
                else:
                    libre = droite - max(l.bbox[2] for l in lignes)
                suivante = rangees[i + 1]
                texte_suivant = "".join(l.text for l in suivante).strip()
                taille = max((s.size for l in lignes for s in l.spans if hasattr(s, "size")), default=10)
                taille_suivante = max((s.size for l in suivante for s in l.spans if hasattr(s, "size")), default=10)
                voulu = (libre > _premier_mot(suivante, est_arabe(texte_suivant)) + 0.3 * taille
                         or bool(docx_post.RE_PUCE.match(texte_suivant))
                         or abs(taille_suivante - taille) > 0.15 * taille)  # un titre suivi de son texte
                if voulu:
                    fin.line_break = 1
                    continue
                # Retour automatique : une espace de couture si la ligne ne finit pas déjà par un blanc ou un trait d'union.
                spans = [s for s in fin.spans if hasattr(s, "chars") and s.chars]
                if spans and not spans[-1].chars[-1].c.isspace() and spans[-1].chars[-1].c not in "-\u00ad":
                    d = spans[-1].chars[-1]
                    spans[-1].chars.append(Char({"c": " ", "bbox": tuple(d.bbox), "origin": d.origin}))
            return None

        self._remplacer(module.RawPageFitz, "_preprocess_text", _preprocess_text)
        self._remplacer(Collection, "sort_in_line_order", sort_in_line_order)
        self._remplacer(Lines, "parse_line_break", parse_line_break)
        return self

    def __exit__(self, *exc):
        for cible, nom, fn in reversed(self.originaux):
            setattr(cible, nom, fn)


def pages_a_ocriser(doc: pymupdf.Document) -> list[int]:
    """Pages sans texte exploitable : aucune lettre, ou trop de caractères inconnus (police sans table Unicode)."""
    pages = []
    for page in doc:
        st = Statistiques()
        lignes = texte_des_lignes(remettre_a_l_endroit(_brut(page), st))
        lettres = sum(1 for _, t, _ in lignes for c in t if c.isalpha())
        if lettres < 15 or st.taux_inconnus() > SEUIL_INCONNUS:
            pages.append(page.number)
    return pages


def elements_par_page(corps) -> list[list]:
    """Le corps d'un Word de pdf2docx, découpé par PAGE : chaque page finit par un paragraphe porteur d'un saut de
    section « nouvelle page » (les sauts « continus » — colonnes — restent dans la page)."""
    from docx.oxml.ns import qn

    pages: list[list] = [[]]
    for el in corps:
        if el.tag == qn("w:sectPr"):
            continue
        pages[-1].append(el)
        sect = el.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") if el.tag == qn("w:p") else None
        if sect is not None:
            typ = sect.find(qn("w:type"))
            if typ is None or typ.get(qn("w:val")) != "continuous":
                pages.append([])
    if not pages[-1]:
        pages.pop()
    return pages


def remplacer_pages_scannees(d, source: Path, scannees: list[int], langues: str, *, vision_dossier: Path | None = None,
                             transcriptions: dict[int, str] | None = None) -> dict:
    """Chaque page scannée : l'image pleine page posée par pdf2docx est remplacée par la page LUE et COMPOSÉE."""
    from docx.oxml.ns import qn

    from . import ocr

    bilan: dict = {}
    corps = d.element.body
    pages = elements_par_page(corps)
    with pymupdf.open(source) as src, tempfile.TemporaryDirectory(prefix="amd-ocr-") as tmp:
        lectures, pngs, largeurs, hauteurs = {}, {}, {}, {}
        for num in scannees:
            png = Path(tmp) / f"p{num}.png"
            pix = src[num].get_pixmap(dpi=ocr.PPP, colorspace=pymupdf.csGRAY)
            pix.save(png)
            pngs[num], largeurs[num], hauteurs[num] = png, pix.width, pix.height
            lectures[num] = ocr.lire_page(png, langues)
        from . import vision

        hf_lus = ocr.entetes_pieds_lus(lectures, hauteurs)
        bilan["_entetes_pieds"] = hf_lus
        a_retirer = [e["texte"] for e in hf_lus["entete"] + hf_lus["pied"]]
        demandes = []
        for num in scannees:
            png, lecture = pngs[num], lectures[num]
            raison = vision.raison_vision(lecture)
            if raison and vision_dossier is not None:
                # Passage 1 : la page est DÉSIGNÉE (image + ce que tesseract a lu) ; l'ERP décide d'appeler ou non.
                vision_dossier.mkdir(parents=True, exist_ok=True)
                cible = vision_dossier / f"page-{num + 1}.png"
                src[num].get_pixmap(dpi=200).save(cible)
                demandes.append({"page": num + 1, "image": str(cible), "raison": raison,
                                 "tesseract": [l["texte"] for l in ocr.lignes_de_mots(lecture["mots"])]})
            texte_vision = (transcriptions or {}).get(num + 1)
            if texte_vision:
                nouveaux, controle = vision.composer_transcription(d, png, lecture, texte_vision, largeurs[num], a_retirer)
                bilan[num + 1] = {"confiance": round(lecture["confiance"], 1), "mots": len(lecture["mots"]), "reconstruite": True,
                                  "figures": len(lecture["figures"]), "tableaux": 0, "vision": True, "raison_vision": raison, "controle": controle}
            else:
                fiable = lecture["mots"] and lecture["confiance"] >= ocr.SEUIL_CONFIANCE
                nouveaux = ocr.composer_dans_word(d, png, lecture, largeurs[num]) if fiable else ocr.composer_image_et_texte(d, png, lecture)
                bilan[num + 1] = {"confiance": round(lecture["confiance"], 1), "mots": len(lecture["mots"]), "reconstruite": bool(fiable),
                                  "figures": len(lecture["figures"]), "tableaux": len(ocr.grille(lecture["traits"])) if fiable else 0,
                                  **({"raison_vision": raison} if raison else {})}
            anciens = pages[num] if num < len(pages) else []
            # Le paragraphe qui porte le saut de section de la page reste (vidé) : il garde taille et marges de la page.
            porteur = next((el for el in anciens if el.tag == qn("w:p") and el.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None), None)
            for el in anciens:
                if el is porteur:
                    for enfant in list(el):
                        if enfant.tag != qn("w:pPr"):
                            el.remove(enfant)
                    continue
                if el.getparent() is not None:
                    corps.remove(el)
            ancre = porteur
            for el in nouveaux:
                corps.remove(el)
            if ancre is not None:
                for el in nouveaux:
                    ancre.addprevious(el)
            else:
                fin = corps.find(qn("w:sectPr"))
                for el in nouveaux:
                    (fin.addprevious(el) if fin is not None else corps.append(el))
        if vision_dossier is not None and demandes:
            import json

            (vision_dossier / "demandes.json").write_text(json.dumps(demandes, ensure_ascii=False), encoding="utf-8")
            bilan["_vision_demandes"] = [x["page"] for x in demandes]
    return bilan


def pdf_vers_docx(source: Path, sortie: Path, *, langues: str = "ara+fra+eng", vision_dossier: Path | None = None,
                  transcriptions: dict[int, str] | None = None) -> dict:
    """Convertit. Rend le bilan (pages, pages OCRisées, paragraphes RTL, en-têtes…) que l'écran et le banc affichent."""
    logging.getLogger().setLevel(logging.ERROR)
    from docx import Document
    from pdf2docx import Converter

    bilan: dict = {"moteur": "pdf2docx+amd", "pages": 0, "pages_ocr": [], "ocr": None}
    with tempfile.TemporaryDirectory(prefix="amd-pdf-docx-") as tmp:
        with pymupdf.open(source) as doc:
            if doc.needs_pass:
                raise ValueError("ce PDF est protégé par un mot de passe")
            bilan["pages"] = doc.page_count
            scannees = pages_a_ocriser(doc)
            hf = reperer_entetes_pieds(doc)
        stats = Statistiques()
        brut = Path(tmp) / "brut.docx"
        with _Patch(hf["zones"], stats):
            cv = Converter(str(source))
            try:
                # Une page qui échoue ne doit pas disparaître en silence (pdf2docx l'ignore par défaut).
                cv.convert(str(brut), multi_processing=False, ignore_page_error=False)
            finally:
                cv.close()
        d = Document(str(brut))
        if scannees:
            bilan["ocr"] = remplacer_pages_scannees(d, source, scannees, langues, vision_dossier=vision_dossier, transcriptions=transcriptions)
            bilan["vision_demandes"] = bilan["ocr"].pop("_vision_demandes", [])
            bilan["pages_ocr"] = [n + 1 for n in scannees]
            lus = bilan["ocr"].pop("_entetes_pieds")
            hf["entete"] = hf["entete"] or lus["entete"]
            hf["pied"] = hf["pied"] or lus["pied"]
        if hf["entete"] or hf["pied"]:
            s0 = d.sections[0]
            if hf["entete"]:
                _poser_zone(s0.header, hf["entete"])
            if hf["pied"]:
                _poser_zone(s0.footer, hf["pied"])
            for s in d.sections[1:]:
                s.header.is_linked_to_previous = True
                s.footer.is_linked_to_previous = True
        bilan["post"] = docx_post.corriger_document(d)
        d.core_properties.title = source.stem
        d.save(str(sortie))
    bilan.update({
        "entetes": [e["texte"] for e in hf["entete"]], "pieds": [e["texte"] for e in hf["pied"]],
        "caracteres": stats.caracteres, "normalises": stats.normalises, "inconnus": stats.inconnus, "lignes_rtl": stats.lignes_rtl,
    })
    return bilan
