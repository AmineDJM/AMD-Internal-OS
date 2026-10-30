"""
L'OCR « VISION » (décision 0019, suite) — quand tesseract ne suffit pas, une transcription par un modèle de vision.

Ce module ne parle à AUCUN service : il dit QUELLES pages méritent la vision (`pages_a_soumettre`), et il COMPOSE une
page à partir d'une transcription reçue (`composer_transcription`), vérifiée contre la lecture de tesseract. L'appel au
modèle, la garde de confidentialité, le plafond de dépense et le journal au nom de la personne vivent côté ERP
(`src/lib/office/conversions-haute-fidelite.ts`, `src/modules/assistant-ia/infra/atelier/ocr-vision.ts`).

Une page est soumise quand :
  • la lecture locale est incertaine (confiance moyenne < `SEUIL_VISION`) ;
  • ou l'arabe y est VOCALISÉ (harakat nombreux, ou mots arabes à basse confiance) — mesuré : tesseract y fait 13 à 58 %
    d'erreurs de caractères, quoi qu'on prépare l'image.

La transcription attendue : le texte de la page, ligne par ligne dans l'ordre de lecture, une ligne vide entre deux
paragraphes, un tableau en lignes « | cellule | cellule | ». La VÉRIFICATION CROISÉE aligne ses lignes sur celles de
tesseract (lettres de base comparées, harakat ignorés) : une ligne que l'un a vue et pas l'autre est COMPTÉE et DITE.
"""

from __future__ import annotations

import io
import re
import statistics
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

from .normaliser import est_arabe, texte_propre

SEUIL_VISION = 80.0
HARAKAT = re.compile(r"[ً-ْٰ]")
LETTRE_ARABE = re.compile(r"[ء-ي]")


def vocalise(texte: str) -> bool:
    """Un texte arabe dont au moins une lettre sur cinq porte une marque : vocalisé."""
    lettres = len(LETTRE_ARABE.findall(texte))
    return lettres >= 20 and len(HARAKAT.findall(texte)) >= 0.2 * lettres


def raison_vision(lecture: dict) -> str | None:
    """Pourquoi cette page mérite la vision — ou `None`."""
    mots = lecture.get("mots") or []
    if not mots:
        return "aucun mot lu par tesseract"
    texte = " ".join(m["t"] for m in mots)
    if vocalise(texte):
        return "arabe vocalisé"
    arabes = [m for m in mots if LETTRE_ARABE.search(m["t"])]
    if arabes and statistics.mean(m["conf"] for m in arabes) < 85 and len(arabes) >= 0.3 * len(mots):
        return "arabe lu avec une confiance basse"
    if lecture.get("confiance", 0) < SEUIL_VISION:
        return f"confiance de lecture {lecture.get('confiance', 0):.0f} %"
    return None


def base(t: str) -> str:
    """Lettres de base seulement (harakat, ponctuation, espaces retirés) — pour aligner deux lectures."""
    t = unicodedata.normalize("NFC", texte_propre(t))
    t = HARAKAT.sub("", t).replace("ـ", "")
    return re.sub(r"[\W_]+", "", t).lower()


def lignes_de_transcription(texte: str) -> list[dict]:
    """La transcription du modèle → lignes typées : texte, ligne vide (fin de paragraphe), rangée de tableau."""
    out = []
    for brute in texte_propre(texte).replace("\r", "").split("\n"):
        l = brute.strip()
        if l.startswith("```"):
            continue
        if not l:
            out.append({"genre": "vide"})
        elif l.startswith("|") and l.endswith("|") and l.count("|") >= 3:
            cellules = [c.strip() for c in l.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cellules if c):
                continue  # la ligne de séparation markdown
            out.append({"genre": "rangee", "cellules": cellules})
        else:
            out.append({"genre": "texte", "texte": re.sub(r"\s{2,}", " ", l)})
    return out


SOUKOUN_ARTICLE = re.compile(r"(?<![\u0621-\u064a])((?:[\u0648\u0641\u0628\u0643][\u064e\u0650]?)?\u0627\u0644)\u0652")


def arbitrer_harakat(texte: str, lignes_tesseract: list[str]) -> str:
    """Le modèle de vision AJOUTE un signe que la page ne porte pas, toujours le même (mesuré sur les trois scans
    vocalisés du corpus) : le soukoun sur le lam de l'article, « الْ » pour « ال ». Il n'est gardé que si la lecture
    locale (tesseract) le voit aussi sur le même mot."""
    vus = {base(m): m for l in lignes_tesseract for m in l.split()}

    def un(m: re.Match) -> str:
        debut = m.start()
        fin = texte.find(" ", debut)
        mot = texte[debut: fin if fin >= 0 else len(texte)]
        lu = vus.get(base(mot))
        if lu is not None and SOUKOUN_ARTICLE.search(lu):
            return m.group(0)
        return m.group(1)

    return SOUKOUN_ARTICLE.sub(un, texte)


def verifier(transcription: list[dict], lignes_tesseract: list[dict]) -> dict:
    """Alignement des lignes (lettres de base) : ce que la vision a vu sans tesseract, et l'inverse."""
    vision = [base(l.get("texte") or " ".join(l.get("cellules", []))) for l in transcription if l["genre"] != "vide"]
    vision = [v for v in vision if len(v) >= 3]
    tess = [base(l["texte"]) for l in lignes_tesseract]
    tess = [t for t in tess if len(t) >= 3]

    def trouve(x: str, dans: list[str]) -> bool:
        return any(x in y or y in x or SequenceMatcher(None, x, y).ratio() >= 0.5 for y in dans)

    absentes_tesseract = [v for v in vision if not trouve(v, tess)]
    absentes_vision = [t for t in tess if not trouve(t, vision)]
    accord = SequenceMatcher(None, "".join(vision), "".join(tess)).ratio() if vision and tess else 0.0
    return {"lignes_vision": len(vision), "lignes_tesseract": len(tess), "absentes_de_tesseract": len(absentes_tesseract),
            "absentes_de_la_vision": len(absentes_vision), "accord_lettres": round(accord, 3)}


def composer_transcription(doc, png: Path, lecture: dict, transcription: str, largeur_px: float, a_retirer: list[str] | None = None) -> tuple[list, dict]:
    """Compose la page à partir de la transcription : paragraphes, tableaux (rangées « | … | »), et les ILLUSTRATIONS
    repérées dans l'image, posées à leur hauteur (la ligne de tesseract la plus proche donne la position)."""
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    from PIL import Image

    from . import ocr

    corps = doc.element.body
    avant = list(corps)
    tess = ocr.lignes_de_mots(lecture["mots"])
    lignes = lignes_de_transcription(arbitrer_harakat(transcription, [l["texte"] for l in tess]))
    # L'en-tête et le pied repérés sur les pages vont dans leurs zones : la transcription ne les répète pas dans le corps.
    retirer = [base(t) for t in (a_retirer or []) if len(base(t)) >= 3]
    lignes = [l for l in lignes if l["genre"] != "texte" or not any(SequenceMatcher(None, base(l["texte"]), r).ratio() >= 0.8 for r in retirer)]
    controle = verifier(lignes, tess)
    h_med = statistics.median(l["h"] for l in tess) if tess else 40
    taille = lambda rtl: Pt(round(max(8.0, min(20.0, h_med * 72 / ocr.PPP * (0.78 if rtl else 0.92))) * 2) / 2)

    # Hauteur de chaque bloc : la ligne de tesseract qui lui ressemble le plus (sinon, la suite de la précédente).
    def hauteur(texte: str, apres: float) -> float:
        b = base(texte)
        meilleure = max(tess, key=lambda l: SequenceMatcher(None, b, base(l["texte"])).ratio(), default=None)
        if meilleure is not None and SequenceMatcher(None, b, base(meilleure["texte"])).ratio() >= 0.5:
            return meilleure["bbox"][1]
        return apres + 1

    blocs: list[tuple[float, str, object]] = []
    y = 0.0
    paragraphe: list[str] = []
    rangees: list[list[str]] = []

    def finir_paragraphe() -> None:
        nonlocal y
        if paragraphe:
            texte = " ".join(paragraphe)
            y = hauteur(paragraphe[0], y)
            blocs.append((y, "p", texte))
            paragraphe.clear()

    def finir_tableau() -> None:
        nonlocal y
        if rangees:
            y = hauteur(" ".join(rangees[0]), y)
            blocs.append((y, "t", [list(r) for r in rangees]))
            rangees.clear()

    for l in lignes:
        if l["genre"] == "rangee":
            finir_paragraphe()
            rangees.append(l["cellules"])
        elif l["genre"] == "vide":
            finir_paragraphe()
            finir_tableau()
        else:
            finir_tableau()
            paragraphe.append(l["texte"])
    finir_paragraphe()
    finir_tableau()
    blocs += [(f[1], "f", f) for f in lecture["figures"]]
    blocs.sort(key=lambda b: b[0])

    with Image.open(png) as im:
        for _, genre, e in blocs:
            if genre == "f":
                b = io.BytesIO()
                im.crop(tuple(int(v) for v in e)).save(b, "PNG")
                doc.add_paragraph().add_run().add_picture(io.BytesIO(b.getvalue()), width=Pt((e[2] - e[0]) * 72 / ocr.PPP))
            elif genre == "p":
                p = doc.add_paragraph()
                r = p.add_run(e)
                rtl = est_arabe(e)
                r.font.size = taille(rtl)
                r.font.name = "Noto Naskh Arabic" if rtl else "Liberation Serif"
            else:
                cols = max(len(r) for r in e)
                rtl = est_arabe(" ".join(" ".join(r) for r in e))
                t = doc.add_table(rows=len(e), cols=cols)
                if "Table Grid" in [s.name for s in doc.styles]:
                    t.style = "Table Grid"
                for i, rangee in enumerate(e):
                    for j in range(cols):
                        t.cell(i, j).paragraphs[0].add_run(rangee[j] if j < len(rangee) else "").font.size = taille(rtl)
                if rtl:
                    bv = OxmlElement("w:bidiVisual")
                    tpr = t._tbl.tblPr
                    style = tpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblStyle")
                    (style.addnext(bv) if style is not None else tpr.insert(0, bv))
    nouveaux = [el for el in corps if el not in avant and el.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr"]
    return nouveaux, controle
