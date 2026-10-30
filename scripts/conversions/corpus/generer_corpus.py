"""
LE CORPUS DES CONVERSIONS (décision 0019) — des PDF dont on connaît le texte exact.

Chaque document est DÉCRIT ici (blocs : titres, paragraphes mis en forme, listes, tableaux, images, en-tête et pied),
composé en Word (python-docx, bidi posé comme le fait Word), puis exporté en PDF par DEUX producteurs différents :
LibreOffice (proche de l'export Word) et Chromium (moteur Skia, autre table ToUnicode). Les « scans » sont ces PDF
rendus en image (200 ppp, niveaux de gris, bruit léger) : il n'y reste plus une lettre, l'OCR doit tout retrouver.

`attendu.json` porte la VÉRITÉ : le texte source dans l'ordre de lecture (en-tête, corps, pied), les paragraphes
arabes, les tableaux (lignes × colonnes), les images. Le banc (`scripts/bench/conversions.ts`) mesure contre lui.

Régénérer (poste de développement — LibreOffice, Chromium de Playwright, PyMuPDF, python-docx, Pillow) :
    python3 scripts/conversions/corpus/generer_corpus.py
Les PDF produits sont VERSIONNÉS : le banc ne dépend pas de la version de LibreOffice du poste.
"""

from __future__ import annotations

import io
import json
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFilter, ImageFont

ICI = Path(__file__).resolve().parent
CHROMIUM = next(iter(sorted(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux64/chrome"))), None)

# ─────────────────────────── Les textes ───────────────────────────

AR_TITRE = "تقرير الشؤون التنظيمية للربع الثالث"
AR_P = [
    "تعمل شركة أدفنتوم فارما على تسجيل الأدوية الجنيسة لدى الوكالة الوطنية للمواد الصيدلانية.",
    "بلغ عدد الملفات المودعة اثني عشر ملفا خلال الفترة من يوليو إلى سبتمبر 2026، منها ثلاثة ملفات في انتظار الرد.",
    "يجب احترام الآجال القانونية المحددة في المرسوم التنفيذي رقم 21-82 لا سيما فيما يخص الإجراءات الاستعجالية.",
    "الصَّيْدَلِيَّةُ المَرْكَزِيَّةُ لِلْمُسْتَشْفَيَاتِ تُطَالِبُ بِشَهَادَةِ المُطَابَقَةِ قَبْلَ كُلِّ تَسْلِيمٍ.",
]
AR_LISTE = ["إعداد ملف التسجيل الكامل", "مراجعة شهادة التحليل للدفعة الأولى", "متابعة لجنة التسعير والتعويض"]
AR_TABLE = [
    ["المنتج", "الشكل الصيدلاني", "الحالة", "السعر (دج)"],
    ["باراسيتامول 500 ملغ", "أقراص", "مسجل", "120,50"],
    ["أموكسيسيلين 1 غ", "أقراص مغلفة", "قيد الدراسة", "345,00"],
    ["إيبوبروفين 400 ملغ", "كبسولات", "مسجل", "210,75"],
]
AR_ENTETE = "أدفنتوم فارما — مديرية الشؤون التنظيمية"
AR_PIED = "وثيقة داخلية — سرية"

FR_TITRE = "Rapport des affaires réglementaires — 3ᵉ trimestre 2026"
FR_P = [
    "Adventum Pharma a déposé douze dossiers d’enregistrement auprès de l’ANPP entre juillet et septembre 2026.",
    "L’efficacité du suivi repose sur la fiche de chaque flacon : numéro de lot, date de péremption et certificat d’analyse.",
    "Les délais réglementaires (décret exécutif n° 21-82) imposent une réponse sous quatre-vingt-dix jours ; au-delà, le dossier est réputé clos.",
    "Cette œuvre collective a mobilisé l’équipe qualité, la pharmacovigilance et le service des achats : 1 250 000,00 DZD engagés, soit 8 % du budget annuel.",
]
FR_LISTE = ["Compléter le module 3 (qualité) du CTD", "Relancer le fournisseur de l’API pour le DMF", "Préparer la réunion du comité des prix"]
FR_TABLE = [
    ["Produit", "Forme", "Statut", "Prix (DZD)"],
    ["Paracétamol 500 mg", "Comprimés", "Enregistré", "120,50"],
    ["Amoxicilline 1 g", "Comprimés pelliculés", "En cours", "345,00"],
    ["Ibuprofène 400 mg", "Gélules", "Enregistré", "210,75"],
]
FR_ENTETE = "Adventum Pharma — Direction des affaires réglementaires"
FR_PIED = "Document interne — confidentiel"

MIXTE = [
    "تم إيداع ملف Amoxicilline 1 g لدى ANPP بتاريخ 15/09/2026 تحت الرقم DZ-2026-114.",
    "Le produit « دواء الضغط » est enregistré sous le n° DZ-2026-115 depuis le 2 octobre.",
    "المسؤول الصيدلاني : Dr Karim Benali — الهاتف 021 45 67 89",
    "La notice bilingue doit reprendre la mention « يحفظ بعيدا عن متناول الأطفال » en tête de page.",
]
COL_FR = [
    "La stratégie de lancement s’appuie sur trois piliers : la disponibilité du produit en officine, la formation des délégués médicaux et la présence aux congrès nationaux.",
    "Le premier pilier exige un stock de sécurité de six semaines chez chaque grossiste répartiteur, vérifié chaque lundi par le service logistique.",
    "Le deuxième pilier repose sur un programme de formation de deux jours, suivi d’une évaluation écrite et d’un accompagnement terrain.",
    "Le troisième pilier concerne les congrès : Alger en novembre, Oran en février et Constantine en avril, avec un stand et un symposium satellite.",
    "Chaque pilier dispose d’un responsable, d’un budget et d’indicateurs suivis mensuellement par la direction générale.",
    "Les résultats du premier trimestre de commercialisation seront présentés au conseil d’administration en janvier 2027.",
]


# ─────────────────────────── Les images ───────────────────────────

def police(taille: int) -> ImageFont.ImageFont:
    for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"):
        if Path(p).exists():
            return ImageFont.truetype(p, taille)
    return ImageFont.load_default()


def image_graphique(largeur: int, hauteur: int, graine: int) -> bytes:
    rnd = random.Random(graine)
    im = Image.new("RGB", (largeur, hauteur), (250, 250, 252))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, largeur - 1, hauteur - 1], outline=(40, 60, 120), width=3)
    n = 6
    for i in range(n):
        h = rnd.randint(hauteur // 6, hauteur - 70)
        x0 = 30 + i * (largeur - 60) // n
        d.rectangle([x0, hauteur - 30 - h, x0 + (largeur - 60) // n - 14, hauteur - 30], fill=(30 + 30 * i, 90, 200 - 25 * i))
    d.text((24, 12), f"T{graine % 4 + 1}", fill=(20, 20, 20), font=police(22))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def image_logo(cote: int) -> bytes:
    im = Image.new("RGB", (cote, cote), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.ellipse([10, 10, cote - 10, cote - 10], fill=(0, 102, 153))
    d.text((cote // 2 - 38, cote // 2 - 22), "AMD", fill=(255, 255, 255), font=police(36))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


# ─────────────────────────── Word, comme Word l'écrit ───────────────────────────

def _rtl_paragraphe(p) -> None:
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        # À SA place dans le schéma (après pStyle/numPr, avant spacing/ind/jc) : LibreOffice ignore un w:bidi mal rangé.
        ppr.insert_element_before(OxmlElement("w:bidi"), "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                                  "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
                                  "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                                  "w:sectPr", "w:pPrChange")


def _run(p, texte: str, *, rtl: bool, gras=False, italique=False, couleur: str | None = None, taille: float | None = None, police_latine="Liberation Serif"):
    r = p.add_run(texte)
    r.bold = gras
    r.italic = italique
    if couleur:
        r.font.color.rgb = RGBColor.from_string(couleur)
    if taille:
        r.font.size = Pt(taille)
    rpr = r._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), police_latine)
    fonts.set(qn("w:hAnsi"), police_latine)
    fonts.set(qn("w:cs"), "Noto Naskh Arabic")
    if rtl:
        rpr.append(OxmlElement("w:rtl"))
        if gras:
            rpr.append(OxmlElement("w:bCs"))
        if taille:
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(int(taille * 2)))
            rpr.append(szcs)
    return r


ARABE = lambda t: sum(1 for c in t if "؀" <= c <= "ۿ") > sum(1 for c in t if c.isascii() and c.isalpha())


def paragraphe(doc, morceaux, *, rtl: bool | None = None, style: str | None = None, apres=6):
    """`morceaux` : [(texte, {gras, italique, couleur, taille})]. La direction suit la majorité des lettres."""
    texte = "".join(m[0] for m in morceaux)
    rtl = ARABE(texte) if rtl is None else rtl
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(apres)
    if rtl:
        _rtl_paragraphe(p)
    for t, f in morceaux:
        _run(p, t, rtl=ARABE(t) or (rtl and not any(c.isascii() and c.isalpha() for c in t)), **f)
    return p


def tableau(doc, lignes, *, rtl: bool):
    t = doc.add_table(rows=len(lignes), cols=len(lignes[0]))
    t.style = "Table Grid"
    if rtl:
        tpr = t._tbl.tblPr
        bv = OxmlElement("w:bidiVisual")
        tpr.append(bv)
    for i, l in enumerate(lignes):
        for j, v in enumerate(l):
            c = t.cell(i, j)
            c.paragraphs[0].text = ""
            p = c.paragraphs[0]
            if rtl:
                _rtl_paragraphe(p)
            _run(p, v, rtl=ARABE(v), gras=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def entete_pied(doc, entete: str, pied: str, rtl: bool):
    s = doc.sections[0]
    for zone, texte in ((s.header, entete), (s.footer, pied)):
        p = zone.paragraphs[0]
        if rtl:
            _rtl_paragraphe(p)
        _run(p, texte, rtl=rtl, italique=True, couleur="555555", taille=9)


def image(doc, octets: bytes, largeur_cm: float):
    doc.add_picture(io.BytesIO(octets), width=Cm(largeur_cm))


# ─────────────────────────── Les documents ───────────────────────────

def doc_fr() -> tuple[Document, dict]:
    doc = Document()
    entete_pied(doc, FR_ENTETE, FR_PIED, rtl=False)
    doc.add_heading(FR_TITRE, level=1)
    paragraphe(doc, [(FR_P[0], {})])
    paragraphe(doc, [("Important : ", {"gras": True, "couleur": "C00000"}), (FR_P[1], {"italique": True})])
    image(doc, image_graphique(640, 360, 1), 12)
    doc.add_heading("Actions prioritaires", level=2)
    for it in FR_LISTE:
        doc.add_paragraph(it, style="List Bullet")
    tableau(doc, FR_TABLE, rtl=False)
    doc.add_page_break()
    doc.add_heading("Délais et engagements", level=2)
    paragraphe(doc, [(FR_P[2], {})])
    paragraphe(doc, [(FR_P[3], {"taille": 13, "couleur": "1F4E79"})])
    image(doc, image_logo(256), 4)
    texte = [FR_ENTETE, FR_TITRE, FR_P[0], "Important : " + FR_P[1], "Actions prioritaires", *FR_LISTE,
             *[c for l in FR_TABLE for c in l], "Délais et engagements", FR_P[2], FR_P[3], FR_PIED]
    return doc, {"texte": texte, "arabes": [], "tableaux": [[4, 4]], "entete": FR_ENTETE, "pied": FR_PIED, "gras": ["Important"], "italique": [FR_P[1][:30]]}


def doc_ar() -> tuple[Document, dict]:
    doc = Document()
    entete_pied(doc, AR_ENTETE, AR_PIED, rtl=True)
    h = doc.add_heading("", level=1)
    _rtl_paragraphe(h)
    _run(h, AR_TITRE, rtl=True, gras=True, taille=16)
    paragraphe(doc, [(AR_P[0], {})])
    paragraphe(doc, [("ملاحظة هامة: ", {"gras": True, "couleur": "C00000"}), (AR_P[1], {})])
    image(doc, image_graphique(480, 320, 2), 10)
    for it in AR_LISTE:
        p = doc.add_paragraph(style="List Bullet")
        _rtl_paragraphe(p)
        _run(p, it, rtl=True)
    tableau(doc, AR_TABLE, rtl=True)
    doc.add_page_break()
    paragraphe(doc, [(AR_P[2], {})])
    paragraphe(doc, [(AR_P[3], {"taille": 14})])
    texte = [AR_ENTETE, AR_TITRE, AR_P[0], "ملاحظة هامة: " + AR_P[1], *AR_LISTE, *[c for l in AR_TABLE for c in l], AR_P[2], AR_P[3], AR_PIED]
    arabes = [AR_TITRE, AR_P[0], "ملاحظة هامة: " + AR_P[1], *AR_LISTE, AR_P[2], AR_P[3]]
    return doc, {"texte": texte, "arabes": arabes, "tableaux": [[4, 4]], "entete": AR_ENTETE, "pied": AR_PIED, "gras": ["ملاحظة هامة"], "italique": []}


def doc_mixte() -> tuple[Document, dict]:
    doc = Document()
    doc.add_heading("Fiche produit — بطاقة المنتج", level=1)
    for t in MIXTE:
        paragraphe(doc, [(t, {})])
    image(doc, image_logo(200), 3.5)
    texte = ["Fiche produit — بطاقة المنتج", *MIXTE]
    return doc, {"texte": texte, "arabes": [MIXTE[0], MIXTE[2]], "tableaux": [], "entete": None, "pied": None, "gras": [], "italique": []}


def doc_colonnes() -> tuple[Document, dict]:
    doc = Document()
    doc.add_heading("Plan de lancement — Amoxicilline 1 g", level=1)
    paragraphe(doc, [("Synthèse destinée au comité de direction.", {"italique": True})])
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    cols = s._sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    c.set(qn("w:num"), "2")
    c.set(qn("w:space"), "567")
    if not cols:
        s._sectPr.append(c)
    for t in COL_FR:
        paragraphe(doc, [(t, {})])
    texte = ["Plan de lancement — Amoxicilline 1 g", "Synthèse destinée au comité de direction.", *COL_FR]
    return doc, {"texte": texte, "arabes": [], "tableaux": [], "entete": None, "pied": None, "gras": [], "italique": ["Synthèse destinée"], "colonnes": 2}


def html_chromium() -> tuple[str, dict]:
    import base64
    img = base64.b64encode(image_graphique(420, 280, 3)).decode()
    lignes = "".join("<tr>" + "".join(f"<{'th' if i == 0 else 'td'}>{c}</{'th' if i == 0 else 'td'}>" for c in l) + "</tr>" for i, l in enumerate(AR_TABLE))
    html = f"""<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><style>
    @page {{ size: A4; margin: 2cm; }}
    body {{ font-family: 'Noto Naskh Arabic', 'Amiri', serif; font-size: 13pt; }}
    h1 {{ font-size: 18pt; color: #1F4E79; }}
    table {{ border-collapse: collapse; margin: 10pt 0; }} td, th {{ border: 1px solid #333; padding: 4pt 8pt; }}
    p.fr {{ direction: ltr; text-align: left; font-family: 'Liberation Serif', serif; }}
    </style></head><body>
    <h1>{AR_TITRE}</h1>
    <p>{AR_P[0]}</p>
    <p><b>ملاحظة هامة: </b>{AR_P[1]}</p>
    <p class="fr">{FR_P[0]}</p>
    <img src="data:image/png;base64,{img}" width="315">
    <table>{lignes}</table>
    <p>{MIXTE[0]}</p>
    </body></html>"""
    texte = [AR_TITRE, AR_P[0], "ملاحظة هامة: " + AR_P[1], FR_P[0], *[c for l in AR_TABLE for c in l], MIXTE[0]]
    return html, {"texte": texte, "arabes": [AR_TITRE, AR_P[0], "ملاحظة هامة: " + AR_P[1], MIXTE[0]], "tableaux": [[4, 4]], "entete": None, "pied": None, "gras": ["ملاحظة هامة"], "italique": []}


# ─────────────────────────── Production ───────────────────────────

def vers_pdf_libreoffice(docx: Path, sortie: Path) -> None:
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{d}/profil", "--headless", "--convert-to", "pdf", "--outdir", d, str(docx)],
                       check=True, capture_output=True, timeout=180)
        shutil.copy(Path(d) / (docx.stem + ".pdf"), sortie)


def vers_pdf_chromium(html: str, sortie: Path) -> None:
    if not CHROMIUM:
        raise SystemExit("Chromium de Playwright introuvable (/opt/pw-browsers)")
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "page.html"
        f.write_text(html, encoding="utf-8")
        subprocess.run([str(CHROMIUM), "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={sortie}", f.as_uri()], check=True, capture_output=True, timeout=120)


def scanner(pdf: Path, sortie: Path, graine: int, ppp: int = 200) -> None:
    """Le PDF rendu en image : plus une lettre dans le fichier. Gris, flou léger, bruit — un scanner de bureau."""
    rnd = random.Random(graine)
    src = fitz.open(pdf)
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=ppp, colorspace=fitz.csGRAY)
        im = Image.frombytes("L", (pix.width, pix.height), pix.samples).filter(ImageFilter.GaussianBlur(0.4))
        px = im.load()
        for _ in range(pix.width * pix.height // 400):
            x, y = rnd.randrange(pix.width), rnd.randrange(pix.height)
            px[x, y] = max(0, px[x, y] - rnd.randint(40, 120))
        b = io.BytesIO()
        im.save(b, "JPEG", quality=80)
        p = out.new_page(width=page.rect.width, height=page.rect.height)
        p.insert_image(p.rect, stream=b.getvalue())
    out.save(sortie, deflate=True)


def main() -> None:
    attendu: dict[str, dict] = {}
    with tempfile.TemporaryDirectory() as d:
        for nom, fab in (("fr-rapport", doc_fr), ("ar-note", doc_ar), ("mixte-ar-fr", doc_mixte), ("fr-deux-colonnes", doc_colonnes)):
            doc, meta = fab()
            docx = Path(d) / f"{nom}.docx"
            doc.save(docx)
            shutil.copy(docx, ICI / f"{nom}.docx")
            vers_pdf_libreoffice(docx, ICI / f"{nom}.pdf")
            attendu[f"{nom}.pdf"] = {**meta, "producteur": "LibreOffice"}
        html, meta = html_chromium()
        vers_pdf_chromium(html, ICI / "ar-chromium.pdf")
        attendu["ar-chromium.pdf"] = {**meta, "producteur": "Chromium"}
    for nom, src, graine in (("scan-fr.pdf", "fr-rapport.pdf", 7), ("scan-ar.pdf", "ar-note.pdf", 11)):
        scanner(ICI / src, ICI / nom, graine)
        attendu[nom] = {**attendu[src], "producteur": "scan 200 ppp", "scan": True}
    for nom, meta in attendu.items():
        with fitz.open(ICI / nom) as pdf:
            meta["pages"] = pdf.page_count
            if not meta.get("scan"):
                imgs = []
                for page in pdf:
                    # Ce qui est DESSINÉ sur la page (LibreOffice partage un dictionnaire de ressources entre les pages).
                    imgs += [[i["width"], i["height"]] for i in page.get_image_info(xrefs=True)]
                meta["images"] = sorted(imgs)
            else:
                meta["images"] = []  # la page entière est une image : sa conservation se mesure autrement
    (ICI / "attendu.json").write_text(json.dumps(attendu, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(json.dumps({k: {"pages": v["pages"], "images": v["images"]} for k, v in attendu.items()}, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
